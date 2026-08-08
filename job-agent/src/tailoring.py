import json
import os
import sqlite3
from pathlib import Path

import anthropic
from docx import Document
from docx.shared import Pt
from dotenv import load_dotenv

from state import DB_PATH, ROOT

load_dotenv(ROOT / "config" / ".env", override=True)

MODEL = "claude-sonnet-5"
MASTER_RESUME_PATH = ROOT / "resume" / "master_resume.json"
OUTPUT_DIR = ROOT / "resume" / "tailored"

TAILOR_PROMPT = """You are tailoring a resume to a job description. You may REORDER bullets,
select which optional bullets to include, and adjust emphasis/wording for keyword alignment.
You must NEVER invent new facts, metrics, or experience not present in the source bullets.

Job description:
{jd_text}

Master resume (JSON, bullets tagged by skill/domain/impact_area):
{resume_json}

Return ONLY a JSON object with this shape:
{{
  "career_summary": [<selected/reordered summary bullet texts, max 3>],
  "experience": [
    {{"company": "...", "title": "...", "start_date": "...", "end_date": "...",
      "bullets": [<selected/reordered bullet texts for this role, keep all bullets but reorder by relevance>]}}
  ],
  "skills_highlight": [<max 10 most relevant skills from the master skills list, ordered by relevance>]
}}
"""


def load_master_resume():
    with open(MASTER_RESUME_PATH) as f:
        return json.load(f)


def strip_code_fence(text):
    text = text.strip()
    if text.startswith("```"):
        text = text.split("\n", 1)[1] if "\n" in text else text[3:]
        if text.rstrip().endswith("```"):
            text = text.rstrip()[:-3]
    return text.strip()


def tailor_resume(client, jd_text, master_resume):
    prompt = TAILOR_PROMPT.format(jd_text=jd_text[:4000], resume_json=json.dumps(master_resume))
    response = client.messages.create(
        model=MODEL,
        max_tokens=8000,
        messages=[{"role": "user", "content": prompt}],
    )
    text = next(b.text for b in response.content if hasattr(b, "text")).strip()
    return json.loads(strip_code_fence(text))


def render_docx(master_resume, tailored, output_path):
    doc = Document()
    doc.add_heading(master_resume["name"], level=0)
    doc.add_paragraph(master_resume["headline"])
    contact = master_resume["contact"]
    doc.add_paragraph(f"{contact['email']} | {contact['phone']} | {contact['location']}")

    doc.add_heading("Career Summary", level=1)
    for bullet in tailored["career_summary"]:
        doc.add_paragraph(bullet, style="List Bullet")

    doc.add_heading("Professional Experience", level=1)
    for role in tailored["experience"]:
        p = doc.add_paragraph()
        run = p.add_run(f"{role['title']}, {role['company']}")
        run.bold = True
        p.add_run(f"  ({role['start_date']} - {role['end_date']})")
        for bullet in role["bullets"]:
            doc.add_paragraph(bullet, style="List Bullet")

    doc.add_heading("Skills", level=1)
    doc.add_paragraph(", ".join(tailored["skills_highlight"]))

    doc.add_heading("Education", level=1)
    for edu in master_resume["education"]:
        doc.add_paragraph(f"{edu['degree']}, {edu['institution']} ({edu['years']}) — {edu['detail']}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def run_tailoring():
    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        raise RuntimeError("ANTHROPIC_API_KEY not set — add it to config/.env")

    client = anthropic.Anthropic(api_key=api_key)
    master_resume = load_master_resume()

    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    rows = conn.execute("SELECT * FROM jobs WHERE status = 'scored'").fetchall()

    tailored_jobs = []
    for row in rows:
        job = dict(row)
        tailored = tailor_resume(client, job["jd_text"] or "", master_resume)
        output_path = OUTPUT_DIR / f"{job['id']}_{job['company']}.docx".replace("/", "_")
        render_docx(master_resume, tailored, output_path)
        conn.execute("UPDATE jobs SET status = 'tailored' WHERE id = ?", (job["id"],))
        conn.commit()
        tailored_jobs.append((job["id"], str(output_path)))

    conn.close()
    return tailored_jobs


if __name__ == "__main__":
    results = run_tailoring()
    print(f"Tailored {len(results)} resumes.")
