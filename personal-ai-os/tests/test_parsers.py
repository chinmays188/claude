import pytest

from app.knowledge.parsers import ParseError, default_parser_registry


def test_text_parser_decodes_utf8():
    registry = default_parser_registry()

    result = registry.parse("notes.txt", "Hello, world.".encode("utf-8"))

    assert result == "Hello, world."


def test_markdown_uses_text_parser():
    registry = default_parser_registry()

    result = registry.parse("notes.md", "# Heading\n\nBody text.".encode("utf-8"))

    assert "Heading" in result


def test_json_parser_pretty_prints():
    registry = default_parser_registry()

    result = registry.parse("data.json", b'{"a": 1, "b": [1, 2]}')

    assert '"a": 1' in result


def test_json_parser_rejects_invalid_json():
    registry = default_parser_registry()

    with pytest.raises(ParseError):
        registry.parse("data.json", b"not json")


def test_csv_parser_joins_rows():
    registry = default_parser_registry()

    result = registry.parse("table.csv", b"a,b,c\n1,2,3\n")

    assert "a, b, c" in result
    assert "1, 2, 3" in result


def test_pdf_parser_extracts_text_from_real_pdf():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    import io
    buf = io.BytesIO()
    writer.write(buf)

    registry = default_parser_registry()
    result = registry.parse("doc.pdf", buf.getvalue())

    assert isinstance(result, str)  # blank page -> empty text, but must not crash


def test_pdf_parser_rejects_corrupt_pdf():
    registry = default_parser_registry()

    with pytest.raises(ParseError):
        registry.parse("doc.pdf", b"not a real pdf")


def test_docx_parser_extracts_paragraphs():
    import docx
    import io

    document = docx.Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")
    buf = io.BytesIO()
    document.save(buf)

    registry = default_parser_registry()
    result = registry.parse("doc.docx", buf.getvalue())

    assert "First paragraph." in result
    assert "Second paragraph." in result


def test_unknown_extension_raises():
    registry = default_parser_registry()

    with pytest.raises(ParseError):
        registry.parse("file.xyz", b"data")
